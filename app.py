import io
import json
import os
import re
import streamlit as st
from docx import Document
from google import genai
from google.genai import types

from config import MODEL_NAME, SYSTEM_INSTRUCTION_2

# Configuração da página Streamlit
st.set_page_config(
    page_title="Agente de Relatórios - Baby Nap",
    page_icon="✏️",
    layout="wide"
)

st.title("🏫 Agente de Relatórios Pedagógicos - Baby Nap")
st.write("Gere relatórios pedagógicos formatados de acordo com a BNCC e a Proposta Antroposófica.")

# Verificação da Chave de API
api_key = os.environ.get("GEMINI_API_KEY")

if not api_key:
    api_key = st.sidebar.text_input("Cole sua GEMINI_API_KEY aqui:", type="password")
    if not api_key:
        st.warning("⚠️ Insira a sua chave API do Google AI Studio para continuar.")
        st.stop()

# Inicializa o cliente Gemini
client = genai.Client(api_key=api_key)

def gerar_relatorio_gemini(rascunho_texto: str) -> str:
    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=rascunho_texto,
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_INSTRUCTION_2,
            temperature=0.3,
            response_mime_type="application/json"
        )
    )
    return response.text

def criar_documento_word(texto_relatorio: str) -> io.BytesIO:
    dados = json.loads(texto_relatorio)
    caminho_modelo = os.path.join(os.path.dirname(__file__), "Modelo_Generico_Baby_Nap.docx")
    doc = Document(caminho_modelo)

    def substituir_marcadores(paragrafo):
        texto = "".join(run.text for run in paragrafo.runs)
        marcadores = list(re.finditer(r"\{\{([A-Z0-9_]+)\}\}", texto))

        for marcador in reversed(marcadores):
            chave = marcador.group(1)
            if chave not in dados:
                continue

            posicoes = []
            posicao = 0
            for run in paragrafo.runs:
                posicoes.append(posicao)
                posicao += len(run.text)

            inicio, fim = marcador.span()
            primeiro = next(
                i for i, run in enumerate(paragrafo.runs)
                if posicoes[i] <= inicio < posicoes[i] + len(run.text)
            )
            ultimo = next(
                i for i, run in enumerate(paragrafo.runs)
                if posicoes[i] <= fim - 1 < posicoes[i] + len(run.text)
            )
            prefixo = paragrafo.runs[primeiro].text[:inicio - posicoes[primeiro]]
            sufixo = paragrafo.runs[ultimo].text[fim - posicoes[ultimo]:]
            valor = str(dados[chave])

            if primeiro == ultimo:
                paragrafo.runs[primeiro].text = prefixo + valor + sufixo
            else:
                paragrafo.runs[primeiro].text = prefixo + valor
                for indice in range(primeiro + 1, ultimo):
                    paragrafo.runs[indice].text = ""
                paragrafo.runs[ultimo].text = sufixo

    def percorrer_paragrafos(container):
        for paragrafo in container.paragraphs:
            yield paragrafo
        for tabela in container.tables:
            for linha in tabela.rows:
                for celula in linha.cells:
                    yield from percorrer_paragrafos(celula)

    for paragrafo in percorrer_paragrafos(doc):
        substituir_marcadores(paragrafo)
    for secao in doc.sections:
        for parte in (secao.header, secao.footer):
            for paragrafo in percorrer_paragrafos(parte):
                substituir_marcadores(paragrafo)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# Formulário de entrada
st.subheader("📥 Observações do Aluno")
rascunho = st.text_area(
    "Cole o rascunho ou as anotações do aluno:",
    height=200,
    placeholder="Exemplo: ANDRÉ. André é uma criança comunicativa, ativa e bastante envolvida..."
)

if st.button("🚀 Gerar Relatório", type="primary"):
    if not rascunho.strip():
        st.error("Por favor, insira as observações do aluno antes de gerar.")
    else:
        with st.spinner("A processar o relatório pedagógico..."):
            try:
                relatorio = gerar_relatorio_gemini(rascunho)
                st.success("Relatório gerado com sucesso!")
                
                st.subheader("📄 Relatório Gerado")
                st.markdown(relatorio)
                
                # Botões de Download
                col1, col2 = st.columns(2)

                nome_aluno = str(json.loads(relatorio).get("ALUNO") or "Aluno").strip()
                nome_aluno = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", nome_aluno)
                nome_aluno = re.sub(r"\s+", "_", nome_aluno).strip("._")
                if not nome_aluno or nome_aluno == "[Preencher]":
                    nome_aluno = "Aluno"
                
                with col1:
                    doc_buffer = criar_documento_word(relatorio)
                    st.download_button(
                        label="📥 Descarregar em Word (.docx)",
                        data=doc_buffer,
                        file_name=f"Relatorio_{nome_aluno}_BabyNap.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
                    
                with col2:
                    st.download_button(
                        label="📄 Descarregar em Texto (.txt)",
                        data=relatorio,
                        file_name="Relatorio_Pedagógico_BabyNap.txt",
                        mime="text/plain"
                    )
            except Exception as e:
                st.error(f"Erro ao gerar relatório: {e}")
